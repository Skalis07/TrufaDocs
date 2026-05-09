"""Smokes end-to-end para las vistas de exportacion."""

import io
import zipfile
from unittest.mock import patch
from xml.etree import ElementTree as etree

from django.test import SimpleTestCase
from docx import Document as DocxDocument


DOCX_EXPORT_URL = "/text/export/docx/"
PDF_EXPORT_URL = "/text/export/pdf/"


def _export_payload() -> dict:
    """Payload minimo pero realista para recorrer parseo y export."""
    return {
        "ui_lang": "en",
        "filename": "qa-cv",
        "doc_font": "Arial",
        "name": "Test Candidate",
        "description": "Profile",
        "email": "test@example.com",
        "phone": "+56 9 5555 5555",
        "linkedin": "linkedin.com/in/test",
        "github": "github.com/test",
        "country": "Chile",
        "city": "Santiago",
        "core_order": "experience,education,skills,extra-1",
        "module_order_map": "",
        "exp_role": ["Engineer"],
        "exp_company": ["ACME"],
        "exp_start": ["2023-01"],
        "exp_end": [""],
        "exp_city": ["Santiago"],
        "exp_country": ["Chile"],
        "exp_tech": ["Python"],
        "exp_highlights": ["Built feature\nLed migration"],
        "edu_degree": ["Computer Science"],
        "edu_institution": ["University X"],
        "edu_start": ["2018-01"],
        "edu_end": ["2022-12"],
        "edu_city": ["Santiago"],
        "edu_country": ["Chile"],
        "edu_items": ["Magna Cum Laude\nResearch scholarship"],
        "skill_category": ["Tools"],
        "skill_items": ["Python, SQL"],
        "extra_section_id": ["extra-1"],
        "extra_title": ["Projects"],
        "extra_mode": ["detailed"],
        "extra_entry_section": ["extra-1"],
        "extra_entry_subtitle": [""],
        "extra_entry_title": ["Platform Project"],
        "extra_entry_where": ["ACME Labs"],
        "extra_entry_tech": ["Django"],
        "extra_entry_start": ["2024-01"],
        "extra_entry_end": ["2024-12"],
        "extra_entry_city": ["Remote"],
        "extra_entry_country": ["Chile"],
        "extra_entry_items_si": [""],
        "extra_entry_items_detailed": ["Shipped v1\nImproved latency"],
    }


def _docx_text(blob: bytes) -> str:
    """Extrae texto del DOCX para asserts de humo."""
    doc = DocxDocument(io.BytesIO(blob))
    chunks: list[str] = []
    chunks.extend(paragraph.text for paragraph in doc.paragraphs if paragraph.text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text = cell.text.strip()
                if text:
                    chunks.append(text)
    return " | ".join(chunks)


def _zip_xml(blob: bytes, name: str) -> bytes:
    """Lee un XML interno del DOCX para asserts sobre fuente/estilos."""
    with zipfile.ZipFile(io.BytesIO(blob)) as zf:
        return zf.read(name)


def _font_names_by_xml(blob: bytes) -> dict[str, set[str]]:
    """Extrae todos los nombres de fuente declarados en los XML del DOCX."""
    fonts: dict[str, set[str]] = {}
    with zipfile.ZipFile(io.BytesIO(blob)) as zf:
        for name in zf.namelist():
            if not name.endswith(".xml"):
                continue
            root = etree.fromstring(zf.read(name))
            found: set[str] = set()
            for elem in root.iter():
                local_tag = elem.tag.split("}")[-1]
                if local_tag == "rFonts":
                    for attr_name, attr_val in elem.attrib.items():
                        local_attr = attr_name.split("}")[-1]
                        if local_attr in {"ascii", "hAnsi", "eastAsia", "cs"}:
                            found.add(str(attr_val))
                elif local_tag in {"fontScheme", "majorFont", "minorFont"}:
                    continue
                elif "}" in elem.tag and elem.tag.split("}")[-1] in {
                    "latin",
                    "ea",
                    "cs",
                    "font",
                } and "typeface" in elem.attrib:
                    found.add(str(elem.attrib["typeface"]))
                elif local_tag == "font":
                    name_attr = next((value for key, value in elem.attrib.items() if key.split("}")[-1] == "name"), None)
                    if name_attr:
                        found.add(str(name_attr))
                elif local_tag == "altName":
                    val_attr = next((value for key, value in elem.attrib.items() if key.split("}")[-1] == "val"), None)
                    if val_attr:
                        found.add(str(val_attr))
            if found:
                fonts[name] = found
    return fonts


class ViewExportSmokeTests(SimpleTestCase):
    """Verifica cableado entre POST estructurado, parser y exportacion."""

    def test_export_docx_returns_a_generated_document_from_structured_post(self) -> None:
        """DOCX debe generarse desde el POST actual sin perder items ni extras."""
        response = self.client.post(DOCX_EXPORT_URL, _export_payload())

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        self.assertIn('filename="qa-cv.docx"', response["Content-Disposition"])

        exported_text = _docx_text(response.content)
        self.assertIn("Engineer", exported_text)
        self.assertIn("Built feature", exported_text)
        self.assertIn("University X", exported_text)
        self.assertIn("Magna Cum Laude", exported_text)
        self.assertIn("PROJECTS", exported_text)
        self.assertIn("Platform Project", exported_text)
        self.assertNotIn("Honors:", exported_text)

    def test_export_pdf_uses_generated_docx_before_conversion(self) -> None:
        """PDF debe generar DOCX intermedio y delegar la conversion final."""
        with patch("editor.views._convert_docx_bytes_to_pdf", return_value=(b"%PDF-1.4 fake", None)) as mocked:
            response = self.client.post(PDF_EXPORT_URL, _export_payload())

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/pdf")
        self.assertIn('filename="qa-cv.pdf"', response["Content-Disposition"])
        self.assertEqual(response.content, b"%PDF-1.4 fake")
        mocked.assert_called_once()
        (docx_bytes,), _ = mocked.call_args
        self.assertTrue(docx_bytes)

    def test_export_docx_forces_selected_font_across_content_and_defaults(self) -> None:
        """La fuente elegida debe quedar aplicada en contenido, estilos, theme y numbering."""
        payload = _export_payload()
        payload["doc_font"] = "STIX Two Text"

        response = self.client.post(DOCX_EXPORT_URL, payload)

        self.assertEqual(response.status_code, 200)
        fonts_by_xml = _font_names_by_xml(response.content)
        self.assertTrue(fonts_by_xml)
        for xml_name, font_names in fonts_by_xml.items():
            self.assertEqual(font_names, {"STIX Two Text"}, msg=xml_name)

    def test_export_docx_english_localizes_current_end_from_post_token(self) -> None:
        """Si el POST trae `Present`, la exportación EN debe renderizar `Present`, no `Actualidad`."""
        payload = _export_payload()
        payload["exp_end"] = ["Present"]
        payload["edu_end"] = ["Present"]

        response = self.client.post(DOCX_EXPORT_URL, payload)

        self.assertEqual(response.status_code, 200)
        exported_text = _docx_text(response.content)
        self.assertIn("Present", exported_text)
        self.assertNotIn("Actualidad", exported_text)
